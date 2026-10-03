"""Run the whole studio on this machine, loopback only, with synthetic demos.

    pip install -r apps/api/requirements.txt -r apps/worker/requirements.txt -r requirements-test.txt
    python scripts/dev_studio.py            # then open http://127.0.0.1:8765

Starts a throwaway Postgres (+pgvector) and Redis from local binaries (or
uses E2E_DATABASE_URL / E2E_REDIS_URL), an S3-compatible store on loopback
standing in for internal MinIO, the real API (dev auth: GROUNDED_ENV=
development) and the real worker. Seeds synthetic briefings so the room has
something to show. Nothing binds beyond 127.0.0.1 and nothing leaves the box.
This is a developer convenience, not a deployment: use infra/compose.yaml.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tests" / "e2e"))

from stack import Stack  # noqa: E402


def workbook() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append(["Northwind Fulfilment — Week 32 Finance WBR"])
    ws.append(["Metric", "This week", "Last week", "Target"])
    rows = [
        ("Net revenue", 4820000, 4510000, 4600000, '"$"#,##0'),
        ("Gross margin", 0.412, 0.405, 0.42, "0.0%"),
        ("Orders shipped", 15200, 14100, 15000, "#,##0"),
        ("On-time delivery", 0.951, 0.962, 0.95, "0.0%"),
        ("Cost per order", 6.84, 7.02, 6.9, '"$"#,##0.00'),
    ]
    for label, now, prior, target, fmt in rows:
        ws.append([label, now, prior, target])
        for col in "BCD":
            ws[f"{col}{ws.max_row}"].number_format = fmt
    risks = wb.create_sheet("Risks")
    risks.append(["Risk", "Primary"])
    risks.append(["West-region carrier capacity is tight through week 34.", "yes"])
    asks = wb.create_sheet("Asks")
    asks.append(["Ask", "Primary"])
    asks.append(["Approve two weekend shifts at the Reno site.", "yes"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def package() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("week-32/finance-pack.xlsx", workbook())
        z.writestr(
            "week-32/ops-memo.md",
            "# Week 32 operations memo\n\nThe Reno site cleared its returns backlog on Tuesday.\n\n"
            "Risk: peak-season hiring is two weeks behind plan.\n\nApprove the overtime budget for week 33.",
        )
    return buf.getvalue()


def memo() -> bytes:
    return (
        "# Billing console launch\n\n## What ships\n\nThe new billing console replaces three legacy screens.\n\n"
        "Finance teams can export invoices in one step.\n\n## Rollout\n\nRollout starts with 12 pilot accounts.\n\n"
        "Approve the pilot list by Friday."
    ).encode()


HANDBOOK = b"""# Welcome to Northwind

## Your first day

Request laptop access from IT on day one.

Collect your badge at the front desk.

Sign in to the HR portal and confirm your bank details.

Book your security training in the first week.

## How we work

Core hours are 10:00 to 15:00.

Post questions in the #help channel.

Join the Monday team sync at 9:30.

Approve your onboarding plan with your manager by Friday.
"""

KNOWLEDGE = b"""# Ledgerly admin guide

## Billing

Invoices are generated on the first business day of each month.

Exports are limited to 10,000 rows per file.

## Members

Only workspace admins can invite teammates.

Admins can reset a member's password from the Members page.

## Support

Email support@ledgerly.internal for access problems.
"""

DEMO_SCRIPT = [
    (0, 3600, "Settings · Workspace · Profile", "Start on the settings page from the top menu."),
    (3600, 7600, "Billing · Download invoices · Export CSV", "Choose Billing to download your invoices or export them as CSV."),
    (7600, 11600, "Members · Invite teammate · Roles", "To add a teammate, open Members and press Invite."),
    (11600, 15000, "Roles · Admin · Viewer", "Pick Viewer if they only need to read reports."),
]


def demo_video(dest: Path) -> Path:
    font = next((str(p) for p in (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),) if p.exists()), None)
    filters = []
    for start, end, screen, _said in DEMO_SCRIPT:
        title, *rest = [part.strip() for part in screen.split("·")]
        sub = "    ".join(rest)
        enable = f"between(t,{start/1000},{end/1000})"
        filters.append(f"drawtext={'fontfile=' + font + ':' if font else ''}text='{title}':fontsize=96:fontcolor=0x1a1814:x=120:y=160:enable='{enable}'")
        filters.append(f"drawtext={'fontfile=' + font + ':' if font else ''}text='{sub}':fontsize=60:fontcolor=0x1d3c34:x=120:y=360:enable='{enable}'")
    total = DEMO_SCRIPT[-1][1] / 1000
    subprocess.run(
        ["ffmpeg", "-y", "-v", "error", "-f", "lavfi", "-i", f"color=c=0xf3f0e8:s=1280x720:d={total}:r=24",
         "-f", "lavfi", "-i", f"sine=frequency=220:duration={total}", "-vf", ",".join(filters),
         "-shortest", "-pix_fmt", "yuv420p", "-c:a", "aac", str(dest)],
        check=True,
    )
    return dest


def http(method: str, url: str, body: bytes | None = None, headers: dict | None = None) -> dict:
    request = urllib.request.Request(url, data=body, method=method, headers=headers or {})
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    with opener.open(request, timeout=60) as response:
        return json.loads(response.read() or b"{}")


def multipart(field_file: tuple[str, bytes, str], kind: str) -> tuple[bytes, str]:
    boundary = "----groundeddev"
    name, data, mime = field_file
    parts = [
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"kind\"\r\n\r\n{kind}\r\n".encode(),
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{name}\"\r\nContent-Type: {mime}\r\n\r\n".encode(),
        data,
        f"\r\n--{boundary}--\r\n".encode(),
    ]
    return b"".join(parts), f"multipart/form-data; boundary={boundary}"


def run_jobs(base: str, briefing_id: str, jobs: tuple[str, ...], title: str) -> bool:
    for job in jobs:
        started = http("POST", f"{base}/api/v1/briefings/{briefing_id}/jobs", json.dumps({"type": job}).encode(), {"content-type": "application/json"})
        for _ in range(600):
            state = http("GET", f"{base}/api/v1/jobs/{started['id']}")
            if state["state"] in {"succeeded", "failed"}:
                break
            time.sleep(0.5)
        if state["state"] != "succeeded":
            print(f"  {title}: {job} failed — {state.get('error')}")
            return False
    return True


def seed_demo(base: str, database_url: str, tmp: Path) -> None:
    """An ask-me demo: synthetic recording + knowledge base. Speech segments are
    synthetic here (no Whisper weights in a dev box); the screen is read by real OCR."""
    from sqlalchemy import create_engine, text

    title = "Ledgerly — billing and members demo"
    created = http("POST", base + "/api/v1/briefings", json.dumps({"title": title, "skill_id": "product-walkthrough"}).encode(), {"content-type": "application/json"})
    video = demo_video(tmp / "ledgerly-demo.mp4")
    for name, data, mime in (("ledgerly-demo.mp4", video.read_bytes(), "video/mp4"), ("ledgerly-admin-guide.md", KNOWLEDGE, "text/markdown")):
        body, ctype = multipart((name, data, mime), "attachment")
        http("POST", f"{base}/api/v1/briefings/{created['id']}/assets", body, {"content-type": ctype})
    if not run_jobs(base, created["id"], ("ingest",), title):
        return
    engine = create_engine(database_url)
    with engine.begin() as conn:
        asset = conn.execute(text("SELECT id FROM source_assets WHERE briefing_id = :b AND kind = 'recording'"), {"b": created["id"]}).scalar_one()
        conn.execute(text("UPDATE source_assets SET duration_ms = :d WHERE id = :a"), {"d": DEMO_SCRIPT[-1][1], "a": asset})
        for start, end, _screen, said in DEMO_SCRIPT:
            conn.execute(text("INSERT INTO transcript_segments (id, asset_id, t_start_ms, t_end_ms, text) VALUES (gen_random_uuid(), :a, :s, :e, :t)"),
                         {"a": asset, "s": start, "e": end, "t": said})
        if shutil.which("tesseract"):
            sys.path.insert(0, str(ROOT / "apps" / "engine"))
            from grounded.providers.local import LocalOCR
            from grounded.screen_text import read_screen

            for seg in read_screen(video, LocalOCR(), interval=1.0, duration_ms=DEMO_SCRIPT[-1][1]):
                conn.execute(text("INSERT INTO transcript_segments (id, asset_id, t_start_ms, t_end_ms, text, speaker) VALUES (gen_random_uuid(), :a, :s, :e, :t, 'screen')"),
                             {"a": asset, "s": seg["t_start_ms"], "e": seg["t_end_ms"], "t": seg["text"]})
    engine.dispose()
    if run_jobs(base, created["id"], ("compile", "render", "index"), title):
        print(f"  seeded: {title}")


def seed(base: str) -> None:
    demos = [
        ("Welcome to Northwind — onboarding", "onboarding-guide", ("northwind-handbook.md", HANDBOOK, "text/markdown")),
        ("Week 32 finance review", "finance-wbr", ("week-32-pack.zip", package(), "application/zip")),
        ("Executive review — Q3", "executive-business-review", ("finance-pack.xlsx", workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ("Billing console launch", "launch-announcement", ("billing-launch.md", memo(), "text/markdown")),
    ]
    for title, skill, upload in demos:
        created = http("POST", base + "/api/v1/briefings", json.dumps({"title": title, "skill_id": skill}).encode(), {"content-type": "application/json"})
        body, ctype = multipart(upload, "attachment")
        http("POST", f"{base}/api/v1/briefings/{created['id']}/assets", body, {"content-type": ctype})
        if run_jobs(base, created["id"], ("ingest", "compile", "render", "index"), title):
            print(f"  seeded: {title}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--no-seed", action="store_true")
    args = parser.parse_args()

    stack = Stack().start()
    tmp = Path(tempfile.mkdtemp(prefix="grounded-dev-"))
    env = {
        **{k: v for k, v in os.environ.items() if not k.startswith(("GROUNDED_", "AWS_", "AUTH_", "OIDC_"))},
        "GROUNDED_DEPLOYMENT_MODE": "offline",
        "GROUNDED_ENV": "development",
        "DEV_BYPASS_AUTH": "true",
        "AUTH_MODE": "dev",
        "DEV_USER_EMAIL": "you@studio.local",
        "DATABASE_URL": stack.database_url,
        "REDIS_URL": stack.redis_url,
        "MINIO_ENDPOINT": stack.s3_url,
        "MINIO_BUCKET": "grounded-dev",
        "MINIO_ACCESS_KEY": "dev",
        "MINIO_SECRET_KEY": "dev",
        "GROUNDED_SKILLS_ROOT": str(ROOT / "skills"),
        "GROUNDED_TTS_PROVIDER": os.environ.get("GROUNDED_TTS_PROVIDER") or ("local" if shutil.which("piper") or os.environ.get("PIPER_BIN") else "none"),
        "GROUNDED_IMAGE_TEXT_PROVIDER": "local-ocr" if shutil.which("tesseract") else "none",
        "ARTIFACT_DIR": str(tmp / "jobs"),
        "HF_HUB_OFFLINE": "1",
        "PYTHONPATH": os.pathsep.join(str(ROOT / p) for p in ("apps/api", "apps/worker", "apps/engine")),
        "NO_PROXY": "127.0.0.1,localhost",
    }
    subprocess.run([sys.executable, "-m", "app.bootstrap"], env=env, check=True, cwd=ROOT / "apps" / "api")
    procs = [
        subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(args.port), "--log-level", "warning"], env=env, cwd=ROOT / "apps" / "api"),
        subprocess.Popen([sys.executable, "-m", "worker.main"], env=env, cwd=ROOT / "apps" / "worker"),
    ]
    base = f"http://127.0.0.1:{args.port}"

    def stop(*_args):
        for proc in procs:
            proc.terminate()
        for proc in procs:
            try:
                proc.wait(5)
            except subprocess.TimeoutExpired:
                proc.kill()
        stack.stop()
        shutil.rmtree(tmp, ignore_errors=True)
        sys.exit(0)

    signal.signal(signal.SIGINT, stop)
    signal.signal(signal.SIGTERM, stop)
    for _ in range(120):
        try:
            http("GET", base + "/health")
            break
        except Exception:  # noqa: BLE001
            time.sleep(0.25)
    if not args.no_seed:
        print("Seeding synthetic briefings…")
        seed(base)
        if shutil.which("ffmpeg"):
            seed_demo(base, stack.database_url, tmp)
    print(f"\nGrounded Studio is running at {base}  (Ctrl+C to stop)\n")
    while all(proc.poll() is None for proc in procs):
        time.sleep(1)
    stop()
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
