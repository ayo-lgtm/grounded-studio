"""Native workbook ingestion for ordinary XLSX and CSV business-review tables."""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Any

from .compile_deck import CompileError

_LABELS = {"metric", "kpi", "line", "name", "measure", "item"}
_ACTUAL = {"actual", "current", "this week", "current week", "value", "result", "h1", "ytd"}
_PRIOR = {"prior", "previous", "last week", "prior week", "previous week", "ly", "prior year"}
_TARGET = {"target", "plan", "budget", "goal", "forecast"}


def parse_workbook(data: bytes) -> dict[str, Any]:
    stripped = data.lstrip()
    if stripped[:1] in {b"{", b"["}:
        try:
            parsed = json.loads(data)
        except json.JSONDecodeError as exc:
            raise CompileError([f"workbook JSON does not parse ({exc.msg})"]) from exc
        if not isinstance(parsed, dict):
            raise CompileError(["workbook JSON must be an object"])
        return parsed
    if data[:8] == _OLE_MAGIC:
        raise CompileError(["legacy .xls workbooks are not supported; save the file as .xlsx"])
    if data[:2] == b"PK":
        return _parse_xlsx(data)
    return _parse_csv(data)


_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class _Book:
    """Formula and cached-value views of one XLSX, with local evaluation."""

    def __init__(self, data: bytes) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CompileError(["XLSX support requires openpyxl"]) from exc
        source = io.BytesIO(data)
        try:
            self.formulas = load_workbook(source, data_only=False, read_only=False)
            source.seek(0)
            self.values = load_workbook(source, data_only=True, read_only=False)
        except Exception as exc:  # noqa: BLE001
            raise CompileError([f"workbook is not a readable XLSX file ({type(exc).__name__})"]) from None
        self.evaluated: set[tuple[str, str]] = set()

    def _authored(self, sheet: str, addr: str) -> Any:
        try:
            return self.formulas[sheet][addr].value
        except KeyError:
            return None

    def value(self, sheet: str, addr: str) -> Any:
        """Cached value, else a locally evaluated simple formula, else None."""
        try:
            cached = self.values[sheet][addr].value
        except KeyError:
            return None
        if cached is not None:
            return cached
        authored = self._authored(sheet, addr)
        if isinstance(authored, str) and authored.startswith("="):
            from .formula import evaluate

            result = evaluate(authored, sheet, self._lookup)
            if result is not None:
                self.evaluated.add((sheet, addr))
            return result
        return authored

    def _lookup(self, sheet: str, addr: str) -> Any:
        try:
            cached = self.values[sheet][addr].value
        except KeyError:
            return None
        return cached if cached is not None else self._authored(sheet, addr)


def _parse_xlsx(data: bytes) -> dict[str, Any]:
    book = _Book(data)
    formulas = book.formulas

    kpis: list[dict[str, Any]] = []
    movers: list[dict[str, Any]] = []
    period = ""
    title = ""
    title_sheet = ""
    title_addr = ""
    period_sheet = ""
    period_addr = ""
    warnings: list[str] = []
    source_ranges: list[dict[str, Any]] = []

    for ws in formulas.worksheets:
        value_ws = _ValueSheet(book, ws.title)
        table = _find_table(ws)
        if table is None:
            continue
        header_row, label_col, actual_col, prior_col, target_col = table
        actual_header = ws.cell(header_row, actual_col).value
        if not period and actual_header is not None:
            period = str(actual_header)
            period_sheet = ws.title
            period_addr = ws.cell(header_row, actual_col).coordinate
        if not title:
            candidate = ws.cell(1, 1).value if header_row > 1 else None
            title = str(candidate).strip() if candidate not in (None, "") else f"{ws.title} Business Review"
            title_sheet = ws.title
            title_addr = ws.cell(1, 1).coordinate if candidate not in (None, "") else ws.cell(header_row, label_col).coordinate

        selected_rows: list[int] = []
        for row in range(header_row + 1, ws.max_row + 1):
            label = ws.cell(row, label_col).value
            if label is None or not str(label).strip():
                continue
            actual_formula = ws.cell(row, actual_col)
            actual_value = value_ws.cell(row, actual_col).value
            if actual_formula.data_type == "f" and actual_value is None:
                warnings.append(
                    f"{ws.title}!{actual_formula.coordinate} has a formula with no cached value "
                    "that cannot be evaluated locally; it is not cited"
                )
                continue
            actual_num = _number(actual_value)
            if actual_num is None:
                continue
            unit = _unit(actual_formula.number_format)
            item: dict[str, Any] = {
                "label": str(label).strip(),
                "sheet": ws.title,
                "addr": actual_formula.coordinate,
                "value": actual_num,
                "display": _display(actual_num, actual_formula.number_format, unit),
                "unit": unit,
                "required": False,
                "formula": actual_formula.value if actual_formula.data_type == "f" else None,
            }
            if target_col is not None:
                target_formula = ws.cell(row, target_col)
                target_value = value_ws.cell(row, target_col).value
                target_num = _number(target_value)
                if target_num is not None:
                    item.update(
                        {
                            "target": target_num,
                            "target_addr": target_formula.coordinate,
                            "target_sheet": ws.title,
                            "target_display": _display(target_num, target_formula.number_format, _unit(target_formula.number_format)),
                            "target_formula": target_formula.value if target_formula.data_type == "f" else None,
                        }
                    )
            kpis.append(item)
            selected_rows.append(row)

            if prior_col is not None:
                prior_formula = ws.cell(row, prior_col)
                prior_value = value_ws.cell(row, prior_col).value
                prior_num = _number(prior_value)
                if prior_num is not None:
                    movers.append(
                        {
                            "label": str(label).strip(),
                            "sheet": ws.title,
                            "addr": actual_formula.coordinate,
                            "value": actual_num,
                            "display": item["display"],
                            "prior_addr": prior_formula.coordinate,
                            "prior": prior_num,
                            "prior_display": _display(prior_num, prior_formula.number_format, _unit(prior_formula.number_format)),
                            "formula": actual_formula.value if actual_formula.data_type == "f" else None,
                            "prior_formula": prior_formula.value if prior_formula.data_type == "f" else None,
                            "unit": unit,
                        }
                    )

        if selected_rows:
            cols = [label_col, actual_col] + [col for col in (prior_col, target_col) if col is not None]
            source_ranges.append(
                _xlsx_source_range(
                    ws,
                    value_ws,
                    header_row,
                    max(selected_rows),
                    min(cols),
                    max(cols),
                )
            )

    if not kpis:
        raise CompileError(
            [
                "No review table found. Use a header row with a metric/name column and an actual/current/value column."
            ]
        )
    first = kpis[0]
    risks, asks = _xlsx_notes(formulas, book)
    for sheet, addr in sorted(book.evaluated):
        warnings.append(f"{sheet}!{addr} had no cached value; evaluated locally from its formula")
    return {
        "title": {"sheet": title_sheet or first["sheet"], "addr": title_addr or first["addr"], "text": title or "Business Review"},
        "period": {"sheet": period_sheet or first["sheet"], "addr": period_addr or first["addr"], "text": period} if period else {},
        "kpis": kpis,
        "movers": movers,
        "risks": risks,
        "asks": asks,
        "data_quality": warnings,
        "source_ranges": source_ranges,
    }


class _ValueSheet:
    """``value_ws.cell(r, c).value`` and ``value_ws[addr].value`` over a _Book."""

    def __init__(self, book: "_Book", title: str) -> None:
        self.book = book
        self.title = title

    def cell(self, row: int, column: int):
        return _Value(self.book.value(self.title, _a1(row, column)))

    def __getitem__(self, addr: str):
        return _Value(self.book.value(self.title, addr))


class _Value:
    def __init__(self, value: Any) -> None:
        self.value = value


_RISK_HEADERS = {"risk", "risks", "key risk", "key risks"}
_ASK_HEADERS = {"ask", "asks", "decision", "decisions", "asks/decisions", "decision needed", "request"}
_PRIMARY_HEADERS = {"primary", "lead", "headline"}


def _xlsx_notes(formulas, book: "_Book") -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Sourced risks and asks: a column headed Risk(s)/Ask(s), or a sheet named so."""
    risks: list[dict[str, Any]] = []
    asks: list[dict[str, Any]] = []
    for ws in formulas.worksheets:
        name = _norm(ws.title)
        sheet_kind = "risk" if name in _RISK_HEADERS else "ask" if name in _ASK_HEADERS else None
        header_row = None
        columns: dict[int, str] = {}
        primary_col = None
        for row in range(1, min(ws.max_row, 30) + 1):
            for col in range(1, ws.max_column + 1):
                header = _norm(ws.cell(row, col).value)
                if header in _RISK_HEADERS:
                    columns[col] = "risk"
                elif header in _ASK_HEADERS:
                    columns[col] = "ask"
                elif header in _PRIMARY_HEADERS:
                    primary_col = col
            if columns:
                header_row = row
                break
        if not columns and sheet_kind:
            header_row, columns = 1, {1: sheet_kind}
            if _norm(ws.cell(1, 1).value) not in _RISK_HEADERS | _ASK_HEADERS:
                header_row = 0
        if not columns:
            continue
        for col, kind in columns.items():
            for row in range(header_row + 1, ws.max_row + 1):
                value = book.value(ws.title, _a1(row, col))
                if value is None or isinstance(value, (int, float)):
                    continue
                text = str(value).strip()
                if not text:
                    continue
                note = {"sheet": ws.title, "addr": _a1(row, col), "text": text}
                if primary_col is not None:
                    flag = _norm(book.value(ws.title, _a1(row, primary_col)))
                    if flag in {"y", "yes", "true", "1", "x", "primary"}:
                        note["primary"] = True
                (risks if kind == "risk" else asks).append(note)
    return risks, asks


def _parse_csv(data: bytes) -> dict[str, Any]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CompileError(["CSV must be UTF-8 encoded"]) from exc
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise CompileError(["CSV is empty"])
    table = _find_csv_table(rows)
    if table is None:
        raise CompileError(["CSV needs a metric/name column and an actual/current/value column"])
    header_row, label_col, actual_col, prior_col, target_col = table
    headers = rows[header_row]
    kpis: list[dict[str, Any]] = []
    movers: list[dict[str, Any]] = []
    selected_rows: list[int] = []
    for rindex, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        if label_col >= len(row) or actual_col >= len(row):
            continue
        label = row[label_col].strip()
        actual = _number(row[actual_col])
        if not label or actual is None:
            continue
        unit = "pct" if "%" in row[actual_col] else "number"
        item: dict[str, Any] = {
            "label": label,
            "sheet": "CSV",
            "addr": _a1(rindex, actual_col + 1),
            "value": actual / 100 if unit == "pct" else actual,
            "display": row[actual_col].strip(),
            "unit": unit,
            "required": False,
        }
        if target_col is not None and target_col < len(row):
            target = _number(row[target_col])
            if target is not None:
                target_unit = "pct" if "%" in row[target_col] else unit
                item.update(
                    {
                        "target": target / 100 if target_unit == "pct" else target,
                        "target_addr": _a1(rindex, target_col + 1),
                        "target_sheet": "CSV",
                        "target_display": row[target_col].strip(),
                    }
                )
        kpis.append(item)
        selected_rows.append(rindex)
        if prior_col is not None and prior_col < len(row):
            prior = _number(row[prior_col])
            if prior is not None:
                prior_value = prior / 100 if "%" in row[prior_col] else prior
                movers.append(
                    {
                        "label": label,
                        "sheet": "CSV",
                        "addr": item["addr"],
                        "value": item["value"],
                        "display": item["display"],
                        "prior_addr": _a1(rindex, prior_col + 1),
                        "prior": prior_value,
                        "prior_display": row[prior_col].strip(),
                        "unit": unit,
                    }
                )
    if not kpis:
        raise CompileError(["CSV review table contains no numeric actual values"])
    title_text = rows[0][0].strip() if header_row > 0 and rows[0] and rows[0][0].strip() else "Business Review"
    cols = [label_col, actual_col] + [col for col in (prior_col, target_col) if col is not None]
    source_ranges = [
        _csv_source_range(rows, header_row + 1, max(selected_rows), min(cols) + 1, max(cols) + 1)
    ] if selected_rows else []
    title_addr = "A1" if header_row > 0 else _a1(header_row + 1, label_col + 1)
    return {
        "title": {"sheet": "CSV", "addr": title_addr, "text": title_text},
        "period": {"sheet": "CSV", "addr": _a1(header_row + 1, actual_col + 1), "text": headers[actual_col]},
        "kpis": kpis,
        "movers": movers,
        "risks": [],
        "asks": [],
        "data_quality": [],
        "source_ranges": source_ranges,
    }



def _xlsx_source_range(ws, value_ws, start_row: int, end_row: int, start_col: int, end_col: int) -> dict[str, Any]:
    rows: list[list[dict[str, Any]]] = []
    for rindex in range(start_row, end_row + 1):
        row: list[dict[str, Any]] = []
        for cindex in range(start_col, end_col + 1):
            authored = ws.cell(rindex, cindex)
            resolved = value_ws.cell(rindex, cindex).value
            value = resolved if resolved is not None else authored.value
            num = _number(resolved)
            unit = _unit(authored.number_format)
            display = _display(num, authored.number_format, unit) if num is not None else str(value or "")
            row.append(
                {
                    "addr": authored.coordinate,
                    "display": display,
                    "value": num,
                    "unit": unit,
                }
            )
        rows.append(row)
    return {
        "sheet": ws.title,
        "range": f"{_a1(start_row, start_col)}:{_a1(end_row, end_col)}",
        "rows": rows,
    }


def _csv_source_range(rows: list[list[str]], start_row: int, end_row: int, start_col: int, end_col: int) -> dict[str, Any]:
    visual_rows: list[list[dict[str, Any]]] = []
    for rindex in range(start_row, end_row + 1):
        source = rows[rindex - 1] if rindex - 1 < len(rows) else []
        visual: list[dict[str, Any]] = []
        for cindex in range(start_col, end_col + 1):
            raw = source[cindex - 1] if cindex - 1 < len(source) else ""
            num = _number(raw)
            unit = "pct" if "%" in raw else "number"
            value = num / 100 if num is not None and unit == "pct" else num
            visual.append(
                {
                    "addr": _a1(rindex, cindex),
                    "display": raw,
                    "value": value,
                    "unit": unit,
                }
            )
        visual_rows.append(visual)
    return {
        "sheet": "CSV",
        "range": f"{_a1(start_row, start_col)}:{_a1(end_row, end_col)}",
        "rows": visual_rows,
    }


def _find_table(ws):
    limit = min(ws.max_row, 30)
    for row in range(1, limit + 1):
        headers = [_norm(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        found = _columns(headers)
        if found:
            label, actual, prior, target = found
            return row, label + 1, actual + 1, None if prior is None else prior + 1, None if target is None else target + 1
    return None


def _find_csv_table(rows):
    for ridx, row in enumerate(rows[:30]):
        found = _columns([_norm(value) for value in row])
        if found:
            label, actual, prior, target = found
            return ridx, label, actual, prior, target
    return None


def _columns(headers: list[str]):
    label = next((i for i, h in enumerate(headers) if h in _LABELS), None)
    if label is None:
        return None
    actual = next((i for i, h in enumerate(headers) if h in _ACTUAL), None)
    prior = next((i for i, h in enumerate(headers) if h in _PRIOR), None)
    target = next((i for i, h in enumerate(headers) if h in _TARGET), None)
    if actual is None:
        candidates = [i for i, h in enumerate(headers) if i != label and h and i not in {prior, target}]
        actual = candidates[-1] if candidates else None
    if actual is None:
        return None
    return label, actual, prior, target


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("$", "").replace("€", "").replace("£", "")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").rstrip("%")
    try:
        value = float(text)
    except ValueError:
        return None
    return -value if negative else value


def _unit(number_format: str | None) -> str:
    fmt = str(number_format or "")
    if "%" in fmt:
        return "pct"
    if any(symbol in fmt for symbol in ("$", "€", "£", "¥")):
        return "currency"
    return "number"


def _display(value: float, number_format: str | None, unit: str) -> str:
    if unit == "pct":
        decimals = 1 if ".0" in str(number_format or "") else 0
        return f"{value * 100:.{decimals}f}%"
    if unit == "currency":
        symbol = next((s for s in ("$", "€", "£", "¥") if s in str(number_format or "")), "$")
        decimals = 2 if ".00" in str(number_format or "") else 0
        return f"{symbol}{value:,.{decimals}f}"
    decimals = 2 if ".00" in str(number_format or "") else (1 if ".0" in str(number_format or "") else 0)
    return f"{value:,.{decimals}f}"


def _a1(row: int, col: int) -> str:
    letters = ""
    n = col
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def extract_workbook_cells(data: bytes) -> list[dict[str, Any]]:
    """Return normalized non-empty cells for persistence and audit/search."""
    if data[:8] == _OLE_MAGIC:
        raise CompileError(["legacy .xls workbooks are not supported; save the file as .xlsx"])
    if data[:2] == b"PK":
        book = _Book(data)
        out: list[dict[str, Any]] = []
        for ws in book.formulas.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    authored = cell.value
                    resolved = book.value(ws.title, cell.coordinate)
                    if authored is None and resolved is None:
                        continue
                    num = _number(resolved)
                    out.append(
                        {
                            "sheet": ws.title,
                            "addr": cell.coordinate,
                            "value_num": num,
                            "value_text": (
                                _display(num, cell.number_format, _unit(cell.number_format))
                                if num is not None
                                else str(resolved if resolved is not None else authored)
                            ),
                            "fmt": str(cell.number_format or ""),
                            "formula": authored if isinstance(authored, str) and authored.startswith("=") else None,
                        }
                    )
        return out

    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CompileError(["CSV must be UTF-8 encoded"]) from exc
    out: list[dict[str, Any]] = []
    for rindex, row in enumerate(csv.reader(io.StringIO(text)), start=1):
        for cindex, value in enumerate(row, start=1):
            if value == "":
                continue
            num = _number(value)
            out.append(
                {
                    "sheet": "CSV",
                    "addr": _a1(rindex, cindex),
                    "value_num": (num / 100 if num is not None and "%" in value else num),
                    "value_text": value,
                    "fmt": "percent" if "%" in value else "",
                    "formula": None,
                }
            )
    return out


def extract_workbook_rows(data: bytes) -> list[dict[str, Any]]:
    """Row-level normalized text for retrieval: "Label - Header: value; ...".

    Each row keeps its sheet and A1 range so a chat answer quoting it cites
    the exact cells. Values are the same displays the deck shows.
    """
    cells = extract_workbook_cells(data)
    by_row: dict[tuple[str, int], list[dict[str, Any]]] = {}
    for cell in cells:
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell["addr"])
        if not match:
            continue
        by_row.setdefault((cell["sheet"], int(match.group(2))), []).append(cell)
    headers: dict[str, dict[str, str]] = {}
    out: list[dict[str, Any]] = []
    for (sheet, row_num), row in sorted(by_row.items(), key=lambda item: (item[0][0], item[0][1])):
        row.sort(key=lambda cell: _col_index(cell["addr"]))
        texts = [str(cell.get("value_text") or "").strip() for cell in row]
        if sheet not in headers and _columns([_norm(text) for text in texts]):
            headers[sheet] = {re.sub(r"\d+", "", cell["addr"]): text for cell, text in zip(row, texts)}
            continue
        names = headers.get(sheet, {})
        parts: list[str] = []
        for cell, text in zip(row, texts):
            if not text:
                continue
            column = re.sub(r"\d+", "", cell["addr"])
            label = names.get(column)
            parts.append(f"{label}: {text}" if label and parts else text)
        if not parts:
            continue
        out.append(
            {
                "sheet": sheet,
                "row_num": row_num,
                "range_ref": f"{row[0]['addr']}:{row[-1]['addr']}" if len(row) > 1 else row[0]["addr"],
                "text": parts[0] + (" - " + "; ".join(parts[1:]) if len(parts) > 1 else ""),
            }
        )
    return out


def _col_index(addr: str) -> int:
    letters = re.match(r"[A-Z]+", addr).group(0)
    total = 0
    for char in letters:
        total = total * 26 + ord(char) - 64
    return total
