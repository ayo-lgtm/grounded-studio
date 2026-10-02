import io
import unittest

from grounded.compile_sources import compile_uploaded
from grounded.ingest_workbook import parse_workbook
from grounded.qa import validate_script


class WorkbookIngestTest(unittest.TestCase):
    def test_csv_business_review_is_discovered(self):
        data = (
            "Metric,Prior,Actual,Target\n"
            "Orders,100,125,120\n"
            "Defect Rate,3%,2%,2.5%\n"
        ).encode()
        pack = parse_workbook(data)
        self.assertEqual(len(pack["kpis"]), 2)
        self.assertEqual(pack["kpis"][0]["label"], "Orders")
        self.assertEqual(pack["kpis"][0]["value"], 125)
        self.assertEqual(pack["kpis"][1]["unit"], "pct")
        self.assertAlmostEqual(pack["kpis"][1]["value"], 0.02)
        self.assertEqual(len(pack["movers"]), 2)

    def test_native_xlsx_preserves_display_and_cells(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")

        wb = Workbook()
        ws = wb.active
        ws.title = "Finance"
        ws.append(["Metric", "Prior", "Actual", "Target"])
        ws.append(["Revenue", 4_000_000, 4_820_000, 4_600_000])
        ws.append(["Margin", 0.39, 0.412, 0.40])
        for cell in ws[2][1:]:
            cell.number_format = '$#,##0'
        for cell in ws[3][1:]:
            cell.number_format = '0.0%'
        buf = io.BytesIO()
        wb.save(buf)

        pack = parse_workbook(buf.getvalue())
        revenue = pack["kpis"][0]
        margin = pack["kpis"][1]
        self.assertEqual(revenue["addr"], "C2")
        self.assertEqual(revenue["display"], "$4,820,000")
        self.assertEqual(margin["display"], "41.2%")
        self.assertEqual(pack["period"]["addr"], "C1")

    def test_finance_skill_compiles_native_csv_with_provenance(self):
        data = (
            "Metric,Prior,Actual,Target\n"
            "Orders,100,125,120\n"
            "Returns,8,6,7\n"
        ).encode()
        script = compile_uploaded(
            skill_id="finance-wbr",
            title="Finance WBR",
            workbook=data,
        )
        self.assertEqual(script["skill_id"], "finance-wbr")
        self.assertEqual(script["renderer"], "deck")
        self.assertEqual(script["provenance"]["skill_id"], "finance-wbr")
        self.assertIn("workbook-analysis", script["provenance"]["craft_versions"])
        self.assertIn("source-range", [beat["layout"] for beat in script["beats"]])
        self.assertEqual(validate_script(script), [])


if __name__ == "__main__":
    unittest.main()
