"""Native ingestion: xlsx, csv, docx, pdf, pptx, txt/md, images, packages."""

import io
import unittest
import zipfile

from openpyxl import Workbook

from grounded.compile_deck import CompileError
from grounded.formula import evaluate
from grounded.ingest import parse_pptx, parse_text
from grounded.ingest_workbook import extract_workbook_cells, parse_workbook
from grounded.sources import SourceError, detect, detect_bytes, normalize, unpack

_SLIDE = (
    '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>{}</p:spTree></p:cSld></p:sld>'
)


def pptx(slides):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/presentation.xml", "<p:presentation/>")
        for index, paragraphs in enumerate(slides, start=1):
            body = "".join(f"<p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>" for text in paragraphs)
            z.writestr(f"ppt/slides/slide{index}.xml", _SLIDE.format(body))
    return buf.getvalue()


def xlsx(rows, sheet="KPI"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def png():
    return b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


class DetectTest(unittest.TestCase):
    def test_magic_bytes(self):
        self.assertEqual(detect("x.pdf", b"%PDF-1.7").kind, "document")
        self.assertEqual(detect("shot.png", png()[:16]).kind, "image")
        self.assertEqual(detect("a.jpg", b"\xff\xd8\xff\xe0").format, "jpeg")
        self.assertEqual(detect("talk.mp4", b"\x00\x00\x00\x18ftypmp42").kind, "recording")
        self.assertEqual(detect("pack.csv", b"Metric,Actual").kind, "workbook")
        self.assertEqual(detect("notes.md", b"# Title").format, "md")

    def test_office_zips_sniffed(self):
        self.assertEqual(detect_bytes("upload.bin", xlsx([["Metric", "Actual"], ["A", 1]])).format, "xlsx")
        self.assertEqual(detect_bytes("deck.zip", pptx([["Title"]])).format, "pptx")

    def test_legacy_binary_office_refused(self):
        with self.assertRaises(SourceError):
            detect("old.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 8)
        with self.assertRaises(CompileError):
            parse_workbook(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)


class PresentationTest(unittest.TestCase):
    def test_slides_become_cited_blocks(self):
        doc = parse_pptx(pptx([["Q3 business review", "Agenda"], ["Revenue grew to $4.8M", "Risk: hiring is behind plan"]]))
        self.assertEqual(doc["title"], "Q3 business review")
        ids = [(b["id"], b["page"]) for b in doc["blocks"]]
        self.assertIn(("s2-b1", 2), ids)
        self.assertEqual(next(b for b in doc["blocks"] if b["id"] == "s2-b2")["role"], "risk")


class TextTest(unittest.TestCase):
    def test_markdown_headings_become_paths(self):
        doc = parse_text(b"# Launch notes\n\n## Scope\n\nWe ship the billing view.\n\nApprove the beta.", markdown=True)
        self.assertEqual(doc["title"], "Launch notes")
        ship = next(b for b in doc["blocks"] if b["text"].startswith("We ship"))
        self.assertEqual(ship["heading_path"], "Scope")
        self.assertEqual(doc["blocks"][-1]["role"], "ask")


class ImageTest(unittest.TestCase):
    def test_screenshot_requires_explicit_provider(self):
        with self.assertRaises(SourceError):
            normalize("shot.png", png())

    def test_screenshot_text_records_extractor(self):
        class OCR:
            name, model_id = "local-ocr", "tesseract"

            def extract_text(self, data, fmt):
                return ["Orders dashboard", "Open tickets 42"]

        result = normalize("shot.png", png(), image_text=OCR())
        self.assertEqual(result.kind, "image")
        self.assertEqual(result.blocks[1]["block_id"], "img-b2")
        self.assertEqual(result.blocks[1]["extractor"], "local-ocr:tesseract")


class PackageTest(unittest.TestCase):
    def _zip(self, members):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in members.items():
                z.writestr(name, data)
        return buf.getvalue()

    def test_mixed_package_unpacks(self):
        data = self._zip({
            "wbr/pack.xlsx": xlsx([["Metric", "Actual"], ["Orders", 10]]),
            "wbr/memo.md": b"# Memo\n\nApprove weekend shifts.",
            "__MACOSX/._pack.xlsx": b"junk",
            ".DS_Store": b"junk",
        })
        members = unpack(data)
        self.assertEqual(sorted(m.detected.kind for m in members), ["document", "workbook"])
        self.assertTrue(all("/" not in m.name for m in members))

    def test_zip_bomb_and_nesting_refused(self):
        with self.assertRaises(SourceError):
            unpack(self._zip({"bomb.txt": b"0" * (10 * 1024 * 1024)}))
        with self.assertRaises(SourceError):
            unpack(self._zip({"inner.zip": self._zip({"a.txt": b"x"})}))


class WorkbookTest(unittest.TestCase):
    def test_formulas_without_cache_are_evaluated_locally(self):
        data = xlsx([["Metric", "Actual", "Target"], ["Orders", 1200, 1150], ["Double", "=B2*2", "=SUM(C2:C2)+10"]])
        pack = parse_workbook(data)
        double = next(k for k in pack["kpis"] if k["label"] == "Double")
        self.assertEqual(double["value"], 2400)
        self.assertEqual(double["target"], 1160)
        self.assertTrue(any("evaluated locally" in w for w in pack["data_quality"]))
        cells = {(c["sheet"], c["addr"]): c for c in extract_workbook_cells(data)}
        self.assertEqual(cells[("KPI", "B3")]["value_num"], 2400)
        self.assertEqual(cells[("KPI", "B3")]["formula"], "=B2*2")

    def test_unsupported_formula_is_not_cited(self):
        data = xlsx([["Metric", "Actual"], ["Orders", 10], ["Lookup", '=VLOOKUP("x",A1:B2,2,FALSE)']])
        pack = parse_workbook(data)
        self.assertNotIn("Lookup", [k["label"] for k in pack["kpis"]])
        self.assertTrue(any("cannot be evaluated" in w for w in pack["data_quality"]))

    def test_formula_grammar_is_closed(self):
        lookup = {("S", "A1"): 2, ("S", "A2"): 3, ("S", "A3"): "=A1+A2"}.get
        self.assertEqual(evaluate("=A1*A2", "S", lambda s, a: lookup((s, a))), 6)
        self.assertEqual(evaluate("=SUM(A1:A3)", "S", lambda s, a: lookup((s, a))), 10)
        self.assertIsNone(evaluate("=INDIRECT(\"A1\")", "S", lambda s, a: lookup((s, a))))
        self.assertIsNone(evaluate("=A1/0", "S", lambda s, a: lookup((s, a))))
        cyclic = {("S", "A1"): "=A2", ("S", "A2"): "=A1"}.get
        self.assertIsNone(evaluate("=A1", "S", lambda s, a: cyclic((s, a))))

    def test_numbers_in_titles_are_cited_to_the_title_cell(self):
        from grounded.dispatch import SourceBundle, compile_bundle

        data = xlsx([["Finance WBR - Week 32"], ["Metric", "Actual"], ["Orders", 10]])
        script = compile_bundle(skill_id="finance-wbr", title="t", bundle=SourceBundle(workbook=data, workbook_asset="w"))
        cover = script["beats"][0]
        self.assertIn({"value": 32, "sheet": "KPI", "addr": "A1", "in_text": True}, cover["claims"])

    def test_numbers_in_workbook_notes_are_cited(self):
        from grounded.dispatch import SourceBundle, compile_bundle

        wb = Workbook()
        ws = wb.active
        ws.title = "KPI"
        ws.append(["Metric", "Actual"])
        ws.append(["Orders", 10])
        risks = wb.create_sheet("Risks")
        risks.append(["Risk"])
        risks.append(["Capacity is tight through week 34."])
        buf = io.BytesIO()
        wb.save(buf)
        script = compile_bundle(skill_id="finance-wbr", title="t", bundle=SourceBundle(workbook=buf.getvalue(), workbook_asset="w"))
        risk = next(b for b in script["beats"] if b["layout"] == "risk")
        self.assertEqual(risk["claims"], [{"value": 34, "sheet": "Risks", "addr": "A2", "in_text": True}])

    def test_risks_and_asks_sheets_are_sourced(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "KPI"
        ws.append(["Metric", "Actual"])
        ws.append(["Orders", 10])
        risks = wb.create_sheet("Risks")
        risks.append(["Risk", "Primary"])
        risks.append(["Carrier capacity is tight.", "yes"])
        buf = io.BytesIO()
        wb.save(buf)
        pack = parse_workbook(buf.getvalue())
        self.assertEqual(pack["risks"], [{"sheet": "Risks", "addr": "A2", "text": "Carrier capacity is tight.", "primary": True}])


if __name__ == "__main__":
    unittest.main()
