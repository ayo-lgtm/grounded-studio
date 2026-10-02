"""End-to-end on a real private stack (offline profile), synthetic data only.

API (OIDC auth) -> private S3 -> worker ingest/compile/render/index ->
persisted citations and claims -> grounded chat -> streamed artifacts.
Also proves: unauthenticated and non-member access is refused, upload caps
stream without buffering, cloud transcription is absent and local
transcription never downloads weights, and no synthetic source content is
written to application logs.
"""

from __future__ import annotations

import io
import json
import logging
import os
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
for sub in ("apps/engine", "apps/api", "apps/worker", "tests/e2e"):
    sys.path.insert(0, str(ROOT / sub))

from stack import Stack  # noqa: E402

MARKER = "ZEPHYR-7Q"  # synthetic token placed in every source; must never reach logs
STACK: Stack | None = None
LOGS = io.StringIO()
KEYS: dict[str, str] = {}
client = None  # TestClient, set in setUpModule
worker = None  # worker.main, set in setUpModule
api = None  # app.main, set in setUpModule


def _keys(tmp: Path) -> None:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from jose import jwk

    private = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    pem = private.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
    public_pem = private.public_key().public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo).decode()
    key = jwk.construct(public_pem, "RS256").to_dict()
    key.update({"kid": "e2e", "use": "sig", "alg": "RS256"})
    (tmp / "jwks.json").write_text(json.dumps({"keys": [key]}))
    KEYS["private"] = pem


def token(sub: str, email: str, ttl: int = 600) -> str:
    from jose import jwt

    now = int(time.time())
    return jwt.encode(
        {"sub": sub, "email": email, "iss": "https://idp.corp.internal", "aud": "grounded-studio", "iat": now, "exp": now + ttl},
        KEYS["private"], algorithm="RS256", headers={"kid": "e2e"},
    )


def setUpModule():
    global STACK, client, worker, api
    STACK = Stack().start()
    tmp = STACK.tmp
    _keys(tmp)
    (tmp / "whisper").mkdir()
    for key in [k for k in os.environ if k.startswith(("GROUNDED_", "AWS_", "MODEL_", "AUTH_", "OIDC_", "DEV_"))]:
        del os.environ[key]
    os.environ.update(
        {
            "GROUNDED_DEPLOYMENT_MODE": "offline",
            "GROUNDED_ENV": "production",
            "DATABASE_URL": STACK.database_url,
            "REDIS_URL": STACK.redis_url,
            "MINIO_ENDPOINT": STACK.s3_url,
            "MINIO_BUCKET": "grounded-e2e",
            "MINIO_ACCESS_KEY": "testing",
            "MINIO_SECRET_KEY": "testing",
            "AUTH_MODE": "oidc",
            "OIDC_ISSUER": "https://idp.corp.internal",
            "OIDC_AUDIENCE": "grounded-studio",
            "OIDC_JWKS_FILE": str(tmp / "jwks.json"),
            "BOOTSTRAP_ADMINS": "analyst@corp.internal",
            "GROUNDED_TTS_PROVIDER": "none",
            "GROUNDED_IMAGE_TEXT_PROVIDER": "none",
            "GROUNDED_MAX_UPLOAD_BYTES": str(3 * 1024 * 1024),
            "WHISPER_CACHE": str(tmp / "whisper"),
            "ARTIFACT_DIR": str(tmp / "jobs"),
            "GROUNDED_SKILLS_ROOT": str(ROOT / "skills"),
            "HF_HUB_OFFLINE": "1",
        }
    )
    handler = logging.StreamHandler(LOGS)
    logging.getLogger().addHandler(handler)
    logging.getLogger().setLevel(logging.DEBUG)

    from app import bootstrap

    bootstrap.main()
    bootstrap.main()  # idempotent second run (migrations re-applied)
    from fastapi.testclient import TestClient

    import app.main as api_module  # startup gates run here
    import worker.main as worker_module

    api = api_module
    worker = worker_module
    client = TestClient(api.app)


def tearDownModule():
    from grounded import guard

    guard.uninstall()
    if STACK is not None:
        STACK.stop()


def auth(sub="u-analyst", email="analyst@corp.internal"):
    return {"Authorization": f"Bearer {token(sub, email)}"}


def drain():
    """Run every queued job through the real worker handler."""
    import redis

    r = redis.from_url(STACK.redis_url)
    while True:
        item = r.rpop("grounded.jobs")
        if not item:
            return
        worker.handle(json.loads(item))


def job(briefing_id, kind, headers=None):
    response = client.post(f"/api/v1/briefings/{briefing_id}/jobs", json={"type": kind}, headers=headers or auth())
    assert response.status_code == 200, response.text
    drain()
    return client.get(f"/api/v1/jobs/{response.json()['id']}", headers=headers or auth()).json()


def workbook_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append([f"{MARKER} Finance WBR"])
    ws.append(["Metric", "This week", "Last week", "Target"])
    ws.append([f"{MARKER} net revenue", 4820000, 4500000, 4600000])
    ws.append(["Fill rate", 0.952, 0.931, 0.96])
    ws.append(["Orders", "=1200+300", 1350, 1400])
    for addr in ("B3", "C3", "D3"):
        ws[addr].number_format = '"$"#,##0'
    for addr in ("B4", "C4", "D4"):
        ws[addr].number_format = "0.0%"
    risks = wb.create_sheet("Risks")
    risks.append(["Risk"])
    risks.append([f"{MARKER} carrier capacity is tight in the west."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def package_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("finance/pack.xlsx", workbook_bytes())
        z.writestr("finance/memo.md", f"# {MARKER} memo\n\nThe Reno warehouse cleared its backlog on Tuesday.\n\nApprove two weekend shifts.".encode())
    return buf.getvalue()


class PrivateStackTest(unittest.TestCase):
    def test_00_health_reports_offline_posture(self):
        body = client.get("/health").json()
        self.assertTrue(body["ok"], body)
        self.assertEqual(body["deployment_mode"], "offline")
        self.assertEqual(body["object_store"], "private")
        self.assertEqual(body["auth"], "oidc")
        self.assertEqual(body["providers"]["inference"], "none")
        room = client.get("/")
        self.assertEqual(room.status_code, 200)
        self.assertIn("What should the room watch?", room.text)
        self.assertIn("default-src 'self'", room.headers["content-security-policy"])
        me = client.get("/api/v1/me", headers=auth()).json()
        self.assertEqual(me["email"], "analyst@corp.internal")
        self.assertEqual(me["role"], "admin")
        self.assertEqual(client.get("/api/v1/me").status_code, 401)

    def test_01_endpoints_fail_closed_without_auth(self):
        for method, path in (
            ("get", "/api/v1/briefings"),
            ("post", "/api/v1/briefings"),
            ("get", "/api/v1/skills"),
            ("post", "/api/v1/help/chat"),
            ("get", "/api/v1/artifacts/00000000-0000-0000-0000-000000000000/file"),
            ("post", "/api/v1/briefings/00000000-0000-0000-0000-000000000000/chat"),
        ):
            response = client.request(method.upper(), path, json={} if method == "post" else None)
            self.assertEqual(response.status_code, 401, path)
        bad = client.get("/api/v1/briefings", headers={"Authorization": "Bearer not-a-jwt"})
        self.assertEqual(bad.status_code, 401)
        expired = client.get("/api/v1/briefings", headers={"Authorization": f"Bearer {token('u', 'u@corp.internal', ttl=-60)}"})
        self.assertEqual(expired.status_code, 401)

    def test_10_finance_wbr_from_a_mixed_package(self):
        created = client.post("/api/v1/briefings", json={"title": "Finance WBR", "skill_id": "finance-wbr"}, headers=auth())
        self.assertEqual(created.status_code, 200, created.text)
        briefing_id = created.json()["id"]
        self.assertEqual(created.json()["skill_version"], "1.1.0")
        upload = client.post(
            f"/api/v1/briefings/{briefing_id}/assets",
            files={"file": ("finance-pack.zip", package_bytes(), "application/zip")},
            data={"kind": "attachment"},
            headers=auth(),
        )
        self.assertEqual(upload.status_code, 200, upload.text)
        self.assertEqual(upload.json()["kind"], "package")
        self.assertEqual(len(upload.json()["sha256"]), 64)

        ingest = job(briefing_id, "ingest")
        self.assertEqual(ingest["state"], "succeeded", ingest)
        compiled = job(briefing_id, "compile")
        self.assertEqual(compiled["state"], "succeeded", compiled)

        detail = client.get(f"/api/v1/briefings/{briefing_id}", headers=auth()).json()
        kinds = sorted(a["kind"] for a in detail["assets"])
        self.assertEqual(kinds, ["document", "package", "workbook"])

        script = client.get(f"/api/v1/briefings/{briefing_id}/script", headers=auth()).json()
        raw = script["raw_json"]
        self.assertEqual(script["skill_id"], "finance-wbr")
        self.assertEqual(script["skill_version"], "1.1.0")
        self.assertIn("workbook-analysis", script["provenance"]["craft_versions"])
        self.assertTrue(script["provenance"]["checks_run"])
        expected = sum(len(beat["citations"]) for beat in raw["beats"])
        self.assertEqual(len(script["citations"]), expected)
        self.assertTrue(any(c["kind"] == "document" for c in script["citations"]))
        layouts = [beat["layout"] for beat in raw["beats"]]
        for layout in ("cover", "versus-target", "movers", "source-range", "risk", "ask"):
            self.assertIn(layout, layouts)

        # Persisted claims reference persisted citations and real cells.
        from sqlalchemy import create_engine, text

        engine = create_engine(STACK.database_url)
        with engine.connect() as conn:
            orphan = conn.execute(text(
                "SELECT count(*) FROM claims c JOIN script_beats b ON b.id = c.beat_id "
                "JOIN script_versions s ON s.id = b.script_id WHERE s.briefing_id = :id AND c.citation_id IS NULL"
            ), {"id": briefing_id}).scalar_one()
            claims = conn.execute(text(
                "SELECT count(*) FROM claims c JOIN script_beats b ON b.id = c.beat_id "
                "JOIN script_versions s ON s.id = b.script_id WHERE s.briefing_id = :id"
            ), {"id": briefing_id}).scalar_one()
            missing = conn.execute(text(
                """
                SELECT count(*) FROM citations c
                JOIN script_beats b ON b.id = c.beat_id JOIN script_versions s ON s.id = b.script_id
                LEFT JOIN workbook_cells w ON w.asset_id = c.asset_id AND w.sheet = c.sheet AND w.addr = c.addr
                WHERE s.briefing_id = :id AND c.kind = 'workbook' AND w.id IS NULL
                """
            ), {"id": briefing_id}).scalar_one()
            evaluated = conn.execute(text(
                "SELECT value_num, formula FROM workbook_cells w JOIN source_assets a ON a.id = w.asset_id "
                "WHERE a.briefing_id = :id AND w.sheet = 'KPI' AND w.addr = 'B5'"
            ), {"id": briefing_id}).first()
        engine.dispose()
        self.assertGreater(claims, 10)
        self.assertEqual(orphan, 0)
        self.assertEqual(missing, 0)
        self.assertEqual(evaluated[0], 1500)
        self.assertEqual(evaluated[1], "=1200+300")

        # Grounded chat.
        answer = client.post(f"/api/v1/briefings/{briefing_id}/chat", json={"question": "What was fill rate this week?"}, headers=auth()).json()
        self.assertFalse(answer["refused"], answer)
        self.assertTrue(answer["citations"])
        self.assertEqual(answer["provider"], "retrieval")
        memo = client.post(f"/api/v1/briefings/{briefing_id}/chat", json={"question": "Which warehouse cleared its backlog?"}, headers=auth()).json()
        self.assertFalse(memo["refused"], memo)
        self.assertEqual(memo["citations"][0]["kind"], "document")
        refused = client.post(f"/api/v1/briefings/{briefing_id}/chat", json={"question": "What is the stock price forecast for 2027?"}, headers=auth()).json()
        self.assertTrue(refused["refused"])
        self.assertEqual(refused["text"], "That is not in this briefing.")

        # Index uses this briefing's sources, local embeddings only.
        indexed = job(briefing_id, "index")
        self.assertEqual(indexed["state"], "succeeded", indexed)
        self.assertEqual(indexed["model_ids"]["embedding"], "local-hash:sha256-bag-768")

        # Render deck + audit sidecars, streamed back through the API.
        rendered = job(briefing_id, "render")
        self.assertEqual(rendered["state"], "succeeded", rendered)
        artifacts = client.get(f"/api/v1/briefings/{briefing_id}/artifacts", headers=auth()).json()["artifacts"]
        kinds = {a["kind"] for a in artifacts}
        for kind in ("deck", "cell-audit", "calculation-audit", "provenance", "quality"):
            self.assertIn(kind, kinds)
        deck = next(a for a in artifacts if a["kind"] == "deck")
        pointer = client.get(f"/api/v1/artifacts/{deck['id']}/content", headers=auth()).json()
        self.assertTrue(pointer["url"].startswith("/api/v1/artifacts/"))
        full = client.get(f"/api/v1/artifacts/{deck['id']}/file", headers=auth())
        self.assertEqual(full.status_code, 200)
        self.assertIn(b"<html", full.content.lower())
        self.assertNotIn(b"https://", full.content)
        part = client.get(f"/api/v1/artifacts/{deck['id']}/file", headers={**auth(), "Range": "bytes=0-9"})
        self.assertEqual(part.status_code, 206)
        self.assertEqual(len(part.content), 10)

        # A signed-in user outside the workspace sees nothing.
        stranger = auth("u-stranger", "stranger@corp.internal")
        self.assertEqual(client.get(f"/api/v1/briefings/{briefing_id}", headers=stranger).status_code, 404)
        self.assertEqual(client.get(f"/api/v1/briefings/{briefing_id}/script", headers=stranger).status_code, 404)
        self.assertEqual(client.get(f"/api/v1/artifacts/{deck['id']}/file", headers=stranger).status_code, 404)
        self.assertEqual(client.post(f"/api/v1/briefings/{briefing_id}/chat", json={"question": "fill rate"}, headers=stranger).status_code, 404)
        self.assertEqual(client.get("/api/v1/briefings", headers=stranger).json()["briefings"], [])

    def test_20_upload_cap_streams_and_cleans_up(self):
        created = client.post("/api/v1/briefings", json={"title": "Cap", "skill_id": "weekly-ops-review"}, headers=auth()).json()
        big = b"%PDF-1.7\n" + os.urandom(3 * 1024 * 1024 + 10)
        response = client.post(
            f"/api/v1/briefings/{created['id']}/assets",
            files={"file": ("big.pdf", big, "application/pdf")},
            headers=auth(),
        )
        self.assertEqual(response.status_code, 413)
        detail = client.get(f"/api/v1/briefings/{created['id']}", headers=auth()).json()
        self.assertEqual(detail["assets"], [])

    def test_30_recording_walkthrough_cites_timestamps(self):
        created = client.post("/api/v1/briefings", json={"title": "Walkthrough", "skill_id": "product-walkthrough"}, headers=auth()).json()
        briefing_id = created["id"]
        with tempfile.TemporaryDirectory() as tmp:
            video = Path(tmp) / "demo.mp4"
            subprocess.run(
                ["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=24:duration=6",
                 "-f", "lavfi", "-i", "sine=frequency=440:duration=6", "-shortest", "-pix_fmt", "yuv420p", str(video)],
                check=True,
            )
            upload = client.post(f"/api/v1/briefings/{briefing_id}/assets", files={"file": ("demo.mp4", video.read_bytes(), "video/mp4")}, headers=auth())
        self.assertEqual(upload.json()["kind"], "recording")

        # Local transcription fails closed: no weights are provisioned and none are downloaded.
        transcribed = job(briefing_id, "transcribe")
        self.assertEqual(transcribed["state"], "failed")
        self.assertTrue("not preloaded" in transcribed["error"] or "not installed" in transcribed["error"], transcribed)
        self.assertEqual(os.listdir(os.environ["WHISPER_CACHE"]), [])

        # Stand in for local faster-whisper output (synthetic segments).
        from sqlalchemy import create_engine, text

        engine = create_engine(STACK.database_url)
        with engine.begin() as conn:
            for start, end, words in ((0, 2000, f"Open the {MARKER} settings page."), (2000, 4500, "Choose Billing to see invoices.")):
                conn.execute(text("INSERT INTO transcript_segments (id, asset_id, t_start_ms, t_end_ms, text) VALUES (gen_random_uuid(), :a, :s, :e, :t)"),
                             {"a": upload.json()["id"], "s": start, "e": end, "t": words})
        engine.dispose()
        compiled = job(briefing_id, "compile")
        self.assertEqual(compiled["state"], "succeeded", compiled)
        script = client.get(f"/api/v1/briefings/{briefing_id}/script", headers=auth()).json()
        recording = [c for c in script["citations"] if c["kind"] == "recording"]
        self.assertTrue(recording)
        self.assertTrue(all(c["t_end_ms"] <= 6100 for c in recording))
        rendered = job(briefing_id, "render")
        self.assertEqual(rendered["state"], "succeeded", rendered)
        kinds = {a["kind"] for a in client.get(f"/api/v1/briefings/{briefing_id}/artifacts", headers=auth()).json()["artifacts"]}
        self.assertIn("video", kinds)
        answer = client.post(f"/api/v1/briefings/{briefing_id}/chat", json={"question": "Where do I see invoices?"}, headers=auth()).json()
        self.assertEqual(answer["citations"][0]["kind"], "recording")

    def test_40_json_workbook_packs_cannot_ship(self):
        created = client.post("/api/v1/briefings", json={"title": "JSON", "skill_id": "weekly-ops-review"}, headers=auth()).json()
        pack = json.dumps({"title": {"text": "x", "sheet": "S", "addr": "A1"}, "kpis": [{"label": "Orders", "sheet": "S", "addr": "B2", "value": 9}]})
        client.post(f"/api/v1/briefings/{created['id']}/assets", files={"file": ("kpi-workbook.json", pack.encode(), "application/json")}, headers=auth())
        compiled = job(created["id"], "compile")
        self.assertEqual(compiled["state"], "failed")
        self.assertIn("not auditable", compiled["error"])

    def test_90_logs_carry_no_source_content(self):
        text = LOGS.getvalue()
        self.assertIn("job.succeeded", text)
        self.assertNotIn(MARKER, text)
        self.assertNotIn("Reno warehouse", text)
        self.assertNotIn("4820000", text.replace(",", ""))


if __name__ == "__main__":
    unittest.main()
